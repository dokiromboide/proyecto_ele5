from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import FileResponse
from sqlalchemy.orm import Session
from sqlalchemy import func
import os, tempfile
from datetime import datetime
from app.core.database import get_db
from app.models.attendance import Attendance, AttendanceSession, AttendanceStatus
from app.models.student import Student
from app.models.user import User
from app.models.course import Course

router = APIRouter(prefix="/reports", tags=["Reportes"])


@router.get("/course/{course_id}/excel")
def export_course_report_excel(course_id: int, db: Session = Depends(get_db)):
    """Exportar reporte de asistencia de un curso en Excel."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment

    course = db.query(Course).filter(Course.id == course_id).first()
    if not course:
        raise HTTPException(status_code=404, detail="Curso no encontrado")

    # Consultar datos
    records = db.query(
        Student.student_code,
        User.full_name,
        AttendanceSession.date,
        AttendanceSession.topic,
        Attendance.status
    ).join(Student, Attendance.student_id == Student.id)\
     .join(User, Student.user_id == User.id)\
     .join(AttendanceSession, Attendance.session_id == AttendanceSession.id)\
     .filter(AttendanceSession.course_id == course_id)\
     .order_by(Student.student_code, AttendanceSession.date).all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Asistencia - {course.code}"

    # Colores Unimayor
    green_fill = PatternFill("solid", fgColor="1B5E20")
    header_font = Font(bold=True, color="FFFFFF")

    # Cabecera
    headers = ["Código", "Estudiante", "Fecha", "Tema", "Estado"]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = green_fill
        cell.alignment = Alignment(horizontal="center")

    # Datos
    for row_idx, record in enumerate(records, 2):
        ws.cell(row=row_idx, column=1, value=record.student_code)
        ws.cell(row=row_idx, column=2, value=record.full_name)
        ws.cell(row=row_idx, column=3, value=record.date.strftime("%Y-%m-%d %H:%M"))
        ws.cell(row=row_idx, column=4, value=record.topic or "")
        ws.cell(row=row_idx, column=5, value=record.status)

    # Auto-ancho columnas
    for col in ws.columns:
        max_length = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[col[0].column_letter].width = min(max_length + 4, 40)

    tmp = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
    wb.save(tmp.name)
    filename = f"asistencia_{course.code}_{datetime.now().strftime('%Y%m%d')}.xlsx"
    return FileResponse(tmp.name, filename=filename, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@router.get("/course/{course_id}/summary")
def get_course_summary(course_id: int, db: Session = Depends(get_db)):
    """Resumen estadístico de asistencia por curso."""
    course = db.query(Course).filter(Course.id == course_id).first()
    if not course:
        raise HTTPException(status_code=404, detail="Curso no encontrado")

    total_sessions = db.query(AttendanceSession).filter(
        AttendanceSession.course_id == course_id
    ).count()

    by_status = db.query(
        Attendance.status, func.count(Attendance.id)
    ).join(AttendanceSession)\
     .filter(AttendanceSession.course_id == course_id)\
     .group_by(Attendance.status).all()

    return {
        "course": {"id": course.id, "name": course.name, "code": course.code},
        "total_sessions": total_sessions,
        "stats": {status: count for status, count in by_status}
    }
